"""
NEXUS CORE OS - CSV & DATA IMPORT API (FASE ONBOARDING KLIEN)
File: apps/core-api/app/api/v1/data_import.py

Modul ini menangani migrasi/impor data massal (bulk import) untuk katalog produk
dan kontak dari sistem legacy seperti Accurate POS/Accounting, Salesforce CRM,
atau file spreadsheet Excel/CSV.

🧠 ATURAN ARSITEKTUR EMAS:
1. HARGA & UANG: Semua nominal harga dikonversi menjadi integer minor units (BIGINT).
   Contoh: "Rp 25.000" atau "25000" -> 25000 (tidak ada desimal/float).
2. TENANCY & RLS: Semua baris data yang diimpor wajib terikat pada `tenant_id`
   yang diekstrak dari header `X-Tenant-ID`. Sesi database diisolasi dengan RLS.
3. FAIL-TOLERANT: Baris yang valid tetap diimpor, baris yang gagal dikumpulkan
   dalam daftar detail error beserta nomor baris dan alasannya.
"""

import io
import re
import csv
import json
from uuid import UUID, uuid4
from typing import List, Dict, Any, Optional
from fastapi import APIRouter, Depends, UploadFile, File, HTTPException, status, Response
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, Field
from sqlalchemy import text
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.exc import IntegrityError, DBAPIError

from app.core.database import get_tenant_db
from app.core.security import get_current_tenant_id

router = APIRouter(prefix="/import", tags=["Data Import (Accurate / Salesforce)"])

# Batas maksimal ukuran file: 5 MB
MAX_FILE_SIZE = 5 * 1024 * 1024

# ==============================================================================
# PYDANTIC RESPONSE SCHEMAS
# ==============================================================================

class RowErrorDetail(BaseModel):
    row: int = Field(..., description="Nomor baris pada file CSV/Excel (1-indexed)")
    identifier: Optional[str] = Field(None, description="SKU atau Nama sebagai identifikasi baris")
    reason: str = Field(..., description="Penyebab kegagalan validasi")

class ImportResultResponse(BaseModel):
    success: bool
    message: str
    total_rows: int
    success_count: int
    error_count: int
    errors: List[RowErrorDetail] = []
    sample_imported: List[Dict[str, Any]] = []

# ==============================================================================
# UTILITY PARSING FUNCTIONS
# ==============================================================================

def clean_header(header: str) -> str:
    """Membersihkan dan menormalisasi string header (lowercase, tanpa spasi/simbol aneh)."""
    if not header:
        return ""
    # Hapus UTF-8 BOM jika ada
    header = header.replace("\ufeff", "").strip().lower()
    return re.sub(r"[\s\-_]+", " ", header)

def parse_rupiah_to_int(val: Any) -> Optional[int]:
    """
    Mengonversi string harga format Indonesia / internasional ke integer minor units (BIGINT).
    Contoh:
    - "Rp 25.000" -> 25000
    - "25.000,00" -> 25000
    - "25000" -> 25000
    - 25000.0 -> 25000
    """
    if val is None:
        return None
    
    if isinstance(val, (int, float)):
        int_val = int(round(val))
        return int_val if int_val >= 0 else None
    
    s = str(val).strip()
    if not s:
        return None
    
    # Hapus simbol mata uang dan whitespace
    s = re.sub(r"(?i)^(rp|idr|\$)\s*", "", s).strip()
    
    # Jika format Indonesia "25.000,00" -> buang desimal ",00"
    if "," in s and "." in s:
        if s.rfind(",") > s.rfind("."):
            # Format Indonesia: titik pemisah ribuan, koma desimal
            s = s[:s.rfind(",")]
            s = s.replace(".", "")
        else:
            # Format US: koma pemisah ribuan, titik desimal
            s = s[:s.rfind(".")]
            s = s.replace(",", "")
    elif "," in s:
        # Cek jika koma dipakai sebagai pemisah ribuan atau desimal
        parts = s.split(",")
        if len(parts) == 2 and len(parts[1]) == 2:
            # Desimal dua digit, contoh: "25000,00"
            s = parts[0].replace(".", "")
        else:
            s = s.replace(",", "")
    elif "." in s:
        # Cek jika titik adalah pemisah ribuan atau desimal
        parts = s.split(".")
        if len(parts) == 2 and len(parts[1]) in (1, 2):
            # Desimal US
            s = parts[0]
        else:
            # Titik ribuan Indonesia
            s = s.replace(".", "")

    # Ambil hanya digit angka
    digits_only = re.sub(r"\D", "", s)
    if not digits_only:
        return None
    
    try:
        res = int(digits_only)
        return res if res >= 0 else None
    except ValueError:
        return None

def parse_stock_to_int(val: Any) -> int:
    """Mengonversi nilai stok ke integer non-negatif (default 0)."""
    if val is None or str(val).strip() == "":
        return 0
    try:
        # Bersihkan format angka jika ada
        cleaned = re.sub(r"[^\d]", "", str(val))
        return int(cleaned) if cleaned else 0
    except Exception:
        return 0

def detect_delimiter(text_sample: str) -> str:
    """Mendeteksi apakah CSV menggunakan koma (,) atau titik-koma (;) khas Excel Indonesia."""
    semicolons = text_sample.count(";")
    commas = text_sample.count(",")
    tabs = text_sample.count("\t")
    
    if semicolons > commas and semicolons > tabs:
        return ";"
    if tabs > commas:
        return "\t"
    return ","

async def read_uploaded_file_rows(file: UploadFile) -> List[Dict[str, str]]:
    """
    Membaca dan mem-parsing baris file yang diunggah (.csv atau .xlsx).
    Mendukung encoding UTF-8, UTF-8-sig (Excel BOM), dan delimiter koma/titik-koma.
    """
    content = await file.read()
    if len(content) > MAX_FILE_SIZE:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Ukuran file melebihi batas maksimal {MAX_FILE_SIZE // (1024 * 1024)}MB."
        )
    
    filename = (file.filename or "").lower()
    rows: List[Dict[str, str]] = []

    # 1. Parsing jika file Excel (.xlsx)
    if filename.endswith(".xlsx"):
        try:
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
            sheet = wb.active
            iter_rows = list(sheet.iter_rows(values_only=True))
            if not iter_rows:
                return []
            
            headers = [clean_header(str(h or "")) for h in iter_rows[0]]
            for r in iter_rows[1:]:
                if not any(r):
                    continue
                row_dict = {}
                for idx, cell in enumerate(r):
                    if idx < len(headers) and headers[idx]:
                        row_dict[headers[idx]] = str(cell).strip() if cell is not None else ""
                rows.append(row_dict)
            return rows
        except ImportError:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Library openpyxl belum terpasang untuk membaca .xlsx. Silakan gunakan format file .csv."
            )
        except Exception as e:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"Gagal memproses file Excel: {str(e)}"
            )

    # 2. Parsing jika file CSV (.csv atau text)
    # Coba decode dengan utf-8-sig (otomatis menangani BOM), fallback ke latin-1
    try:
        decoded_text = content.decode("utf-8-sig")
    except UnicodeDecodeError:
        try:
            decoded_text = content.decode("latin-1")
        except Exception as e:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"Gagal mendeteksi encoding file teks: {str(e)}"
            )

    sample = decoded_text[:2048]
    delimiter = detect_delimiter(sample)

    reader = csv.reader(io.StringIO(decoded_text), delimiter=delimiter)
    try:
        header_row = next(reader)
    except StopIteration:
        return []

    headers = [clean_header(h) for h in header_row]

    for r in reader:
        if not r or not any(field.strip() for field in r):
            continue
        row_dict = {}
        for idx, val in enumerate(r):
            if idx < len(headers) and headers[idx]:
                row_dict[headers[idx]] = val.strip()
        rows.append(row_dict)

    return rows

# ==============================================================================
# ENDPOINT 1: IMPORT ITEMS (PRODUK & JASA DARI ACCURATE / ERP)
# ==============================================================================

@router.post(
    "/items",
    response_model=ImportResultResponse,
    status_code=status.HTTP_200_OK,
    summary="Import data produk & jasa dari file CSV/Excel Accurate"
)
async def import_items(
    file: UploadFile = File(...),
    db: AsyncSession = Depends(get_tenant_db),
    current_tenant_id: UUID = Depends(get_current_tenant_id)
):
    """
    Mengimpor data katalog massal ke tabel `public.items` dan `public.stock_levels`.
    
    Mapping Kolom Otomatis:
    - SKU: 'sku', 'kode', 'kode barang', 'item code'
    - Nama: 'nama barang', 'nama produk', 'nama', 'name', 'deskripsi'
    - Harga Jual: 'harga jual', 'harga', 'base_price', 'price', 'unit price'
    - Stok Awal: 'stok', 'stock', 'qty', 'jumlah', 'quantity'
    - Tipe Item: 'tipe', 'type', 'item_type', 'jenis' ('stocked'/'service')
    - Kategori/Brand: otomatis disimpan di JSONB `attributes`
    """
    raw_rows = await read_uploaded_file_rows(file)
    if not raw_rows:
        return ImportResultResponse(
            success=False,
            message="File tidak memiliki data atau baris kosong.",
            total_rows=0,
            success_count=0,
            error_count=0,
            errors=[]
        )

    # Ambil SKU yang sudah terdaftar di tenant aktif untuk mencegah duplikasi
    existing_skus_res = await db.execute(
        text("SELECT sku FROM public.items WHERE tenant_id = :tenant_id"),
        {"tenant_id": str(current_tenant_id)}
    )
    existing_db_skus = {row[0].strip().upper() for row in existing_skus_res.fetchall()}

    valid_items_to_insert: List[Dict[str, Any]] = []
    errors: List[RowErrorDetail] = []
    seen_file_skus: set = set()

    for row_idx, row in enumerate(raw_rows, start=2): # Baris 1 adalah header
        # 1. Ekstraksi SKU
        sku_val = (
            row.get("sku") or 
            row.get("kode") or 
            row.get("kode barang") or 
            row.get("item code") or 
            ""
        ).strip()

        # 2. Ekstraksi Nama
        name_val = (
            row.get("nama barang") or 
            row.get("nama produk") or 
            row.get("nama") or 
            row.get("name") or 
            row.get("deskripsi") or 
            ""
        ).strip()

        if not name_val:
            errors.append(RowErrorDetail(
                row=row_idx,
                identifier=sku_val or f"Baris {row_idx}",
                reason="Nama barang/produk wajib diisi."
            ))
            continue

        # Jika SKU kosong, generate SKU otomatis
        if not sku_val:
            sku_val = f"ITEM-{uuid4().hex[:8].upper()}"

        sku_upper = sku_val.upper()

        # Cek duplikat dalam file yang sama
        if sku_upper in seen_file_skus:
            errors.append(RowErrorDetail(
                row=row_idx,
                identifier=sku_val,
                reason=f"SKU '{sku_val}' duplikat di dalam file ini."
            ))
            continue
        
        # Cek duplikat dengan database
        if sku_upper in existing_db_skus:
            errors.append(RowErrorDetail(
                row=row_idx,
                identifier=sku_val,
                reason=f"SKU '{sku_val}' sudah terdaftar di sistem."
            ))
            continue

        seen_file_skus.add(sku_upper)

        # 3. Ekstraksi Harga Jual (Wajib Integer Minor Units)
        raw_price = (
            row.get("harga jual") or 
            row.get("harga") or 
            row.get("base price") or 
            row.get("price") or 
            row.get("unit price") or 
            None
        )
        base_price = parse_rupiah_to_int(raw_price)
        if base_price is None:
            errors.append(RowErrorDetail(
                row=row_idx,
                identifier=sku_val,
                reason=f"Harga jual tidak valid atau bernilai negatif ('{raw_price}')."
            ))
            continue

        # 4. Ekstraksi Stok
        raw_stock = (
            row.get("stok") or 
            row.get("stock") or 
            row.get("qty") or 
            row.get("jumlah") or 
            row.get("quantity") or 
            0
        )
        stock_qty = parse_stock_to_int(raw_stock)

        # 5. Ekstraksi Tipe Item
        raw_type = (
            row.get("tipe") or 
            row.get("type") or 
            row.get("item type") or 
            row.get("jenis") or 
            "stocked"
        ).strip().lower()

        if any(keyword in raw_type for keyword in ["jasa", "service", "layanan"]):
            item_type = "service"
        elif "digital" in raw_type:
            item_type = "digital"
        else:
            item_type = "stocked"

        # 6. Kumpulkan atribut tambahan ke JSONB
        category_val = row.get("kategori") or row.get("category") or ("Jasa" if item_type == "service" else "Sparepart")
        brand_val = row.get("brand") or row.get("merek") or ""
        warranty_val = row.get("garansi") or row.get("warranty") or ""
        
        attributes = {
            "source": "accurate_csv_import",
            "category": category_val,
            "stock": stock_qty
        }
        if brand_val:
            attributes["brand"] = brand_val
        if warranty_val:
            attributes["warranty"] = warranty_val

        # Simpan kolom ekstra lainnya yang tidak termasuk standar
        for k, v in row.items():
            if k not in ["sku", "nama barang", "nama produk", "nama", "harga jual", "harga", "stok", "tipe", "kategori", "brand"]:
                if v:
                    attributes[k] = v

        valid_items_to_insert.append({
            "id": str(uuid4()),
            "tenant_id": str(current_tenant_id),
            "sku": sku_val,
            "name": name_val,
            "item_type": item_type,
            "base_price": base_price,
            "attributes": json.dumps(attributes),
            "stock_qty": stock_qty
        })

    # Bulk Insert ke tabel items & stock_levels
    sample_imported: List[Dict[str, Any]] = []
    
    if valid_items_to_insert:
        try:
            for item in valid_items_to_insert:
                # Insert item
                await db.execute(
                    text("""
                        INSERT INTO public.items (
                            id, tenant_id, sku, name, item_type, base_price, attributes, active
                        ) VALUES (
                            :id, :tenant_id, :sku, :name, :item_type, :base_price, :attributes::jsonb, true
                        )
                    """),
                    {
                        "id": item["id"],
                        "tenant_id": item["tenant_id"],
                        "sku": item["sku"],
                        "name": item["name"],
                        "item_type": item["item_type"],
                        "base_price": item["base_price"],
                        "attributes": item["attributes"]
                    }
                )

                # Jika item fisik dan memiliki stok awal > 0, inisialisasi di stock_levels
                if item["item_type"] == "stocked" and item["stock_qty"] > 0:
                    await db.execute(
                        text("""
                            INSERT INTO public.stock_levels (
                                id, tenant_id, item_id, on_hand, reserved
                            ) VALUES (
                                gen_random_uuid(), :tenant_id, :item_id, :on_hand, 0
                            )
                            ON CONFLICT (tenant_id, item_id) 
                            DO UPDATE SET on_hand = public.stock_levels.on_hand + EXCLUDED.on_hand
                        """),
                        {
                            "tenant_id": item["tenant_id"],
                            "item_id": item["id"],
                            "on_hand": item["stock_qty"]
                        }
                    )
                    
                    # Catat riwayat di stock_movements
                    await db.execute(
                        text("""
                            INSERT INTO public.stock_movements (
                                id, tenant_id, item_id, qty_delta, reason, ref_id
                            ) VALUES (
                                gen_random_uuid(), :tenant_id, :item_id, :qty_delta, 'initial_import', 'ACCURATE-CSV'
                            )
                        """),
                        {
                            "tenant_id": item["tenant_id"],
                            "item_id": item["id"],
                            "qty_delta": item["stock_qty"]
                        }
                    )

            await db.commit()

            sample_imported = [
                {
                    "sku": it["sku"],
                    "name": it["name"],
                    "base_price": it["base_price"],
                    "item_type": it["item_type"],
                    "stock": it["stock_qty"]
                }
                for it in valid_items_to_insert[:5]
            ]
        except (IntegrityError, DBAPIError) as e:
            await db.rollback()
            raise HTTPException(
                status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
                detail=f"Terjadi kesalahan saat menyimpan data massal: {str(e)}"
            )

    return ImportResultResponse(
        success=len(valid_items_to_insert) > 0,
        message=f"Berhasil mengimpor {len(valid_items_to_insert)} produk. {len(errors)} baris dilewati/gagal.",
        total_rows=len(raw_rows),
        success_count=len(valid_items_to_insert),
        error_count=len(errors),
        errors=errors,
        sample_imported=sample_imported
    )

# ==============================================================================
# ENDPOINT 2: IMPORT CONTACTS (KONTAK & CRM DARI SALESFORCE / ACCURATE)
# ==============================================================================

@router.post(
    "/contacts",
    response_model=ImportResultResponse,
    status_code=status.HTTP_200_OK,
    summary="Import data kontak CRM dari file CSV/Excel Salesforce"
)
async def import_contacts(
    file: UploadFile = File(...),
    db: AsyncSession = Depends(get_tenant_db),
    current_tenant_id: UUID = Depends(get_current_tenant_id)
):
    """
    Mengimpor data kontak CRM ke tabel `public.contacts`.
    
    Mapping Kolom Otomatis:
    - Nama: 'nama', 'name', 'display_name', 'nama kontak', 'customer'
    - No HP: 'no hp', 'no. hp', 'phone', 'telepon', 'whatsapp', 'nomor hp'
    - Email: 'email', 'e-mail', 'surel'
    - Perusahaan: 'perusahaan', 'company', 'company_name', 'instansi'
    - Tipe/Jenis: 'tipe', 'kind' ('person'/'org', default 'person')
    """
    raw_rows = await read_uploaded_file_rows(file)
    if not raw_rows:
        return ImportResultResponse(
            success=False,
            message="File tidak memiliki data atau baris kosong.",
            total_rows=0,
            success_count=0,
            error_count=0,
            errors=[]
        )

    valid_contacts: List[Dict[str, Any]] = []
    errors: List[RowErrorDetail] = []

    for row_idx, row in enumerate(raw_rows, start=2):
        name_val = (
            row.get("nama") or 
            row.get("name") or 
            row.get("display name") or 
            row.get("nama kontak") or 
            row.get("customer") or 
            row.get("pelanggan") or 
            ""
        ).strip()

        if not name_val:
            errors.append(RowErrorDetail(
                row=row_idx,
                identifier=f"Baris {row_idx}",
                reason="Nama kontak/pelanggan tidak boleh kosong."
            ))
            continue

        phone_val = (
            row.get("no hp") or 
            row.get("no. hp") or 
            row.get("phone") or 
            row.get("telepon") or 
            row.get("whatsapp") or 
            row.get("nomor hp") or 
            ""
        ).strip()

        email_val = (
            row.get("email") or 
            row.get("e mail") or 
            row.get("surel") or 
            ""
        ).strip()

        company_val = (
            row.get("perusahaan") or 
            row.get("company") or 
            row.get("company name") or 
            row.get("instansi") or 
            ""
        ).strip()

        raw_kind = (
            row.get("tipe") or 
            row.get("kind") or 
            row.get("jenis") or 
            ""
        ).strip().lower()

        kind = "org" if ("perusahaan" in raw_kind or "org" in raw_kind or "pt" in raw_kind or company_val and not phone_val) else "person"

        attributes = {
            "source": "salesforce_csv_import"
        }
        if phone_val:
            attributes["phone"] = phone_val
        if email_val:
            attributes["email"] = email_val
        if company_val:
            attributes["company_name"] = company_val

        # Ekstraksi atribut lainnya
        for k, v in row.items():
            if k not in ["nama", "name", "no hp", "phone", "email", "perusahaan", "company", "tipe", "kind"]:
                if v:
                    attributes[k] = v

        valid_contacts.append({
            "id": str(uuid4()),
            "tenant_id": str(current_tenant_id),
            "display_name": name_val,
            "kind": kind,
            "attributes": json.dumps(attributes)
        })

    sample_imported: List[Dict[str, Any]] = []

    if valid_contacts:
        try:
            for c in valid_contacts:
                await db.execute(
                    text("""
                        INSERT INTO public.contacts (
                            id, tenant_id, display_name, kind, attributes
                        ) VALUES (
                            :id, :tenant_id, :display_name, :kind, :attributes::jsonb
                        )
                    """),
                    {
                        "id": c["id"],
                        "tenant_id": c["tenant_id"],
                        "display_name": c["display_name"],
                        "kind": c["kind"],
                        "attributes": c["attributes"]
                    }
                )
            await db.commit()

            sample_imported = [
                {
                    "display_name": c["display_name"],
                    "kind": c["kind"],
                    "attributes": json.loads(c["attributes"])
                }
                for c in valid_contacts[:5]
            ]
        except (IntegrityError, DBAPIError) as e:
            await db.rollback()
            raise HTTPException(
                status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
                detail=f"Terjadi kesalahan saat menyimpan data kontak: {str(e)}"
            )

    return ImportResultResponse(
        success=len(valid_contacts) > 0,
        message=f"Berhasil mengimpor {len(valid_contacts)} kontak. {len(errors)} baris dilewati/gagal.",
        total_rows=len(raw_rows),
        success_count=len(valid_contacts),
        error_count=len(errors),
        errors=errors,
        sample_imported=sample_imported
    )

# ==============================================================================
# ENDPOINT 3: DOWNLOAD TEMPLATE CSV
# ==============================================================================

@router.get("/template/items", summary="Download contoh template CSV produk Accurate")
async def download_items_template():
    """Mengunduh file template CSV siap pakai untuk import produk."""
    csv_content = (
        "SKU,Nama Barang,Harga Jual,Stok,Tipe,Kategori,Brand\n"
        "LCD-IPH-11,LCD iPhone 11 Original OEM,650000,15,stocked,Sparepart,Apple\n"
        "BAT-SAM-S21,Baterai Samsung Galaxy S21,320000,10,stocked,Sparepart,Samsung\n"
        "CAS-MAG-CLR,Casing MagSafe Clear Armor,75000,50,stocked,Aksesoris,Universal\n"
        "SRV-PASANG-LCD,Jasa Pasang LCD & Kalibrasi,100000,0,service,Jasa Service,Internal\n"
    )
    return Response(
        content=csv_content,
        media_type="text/csv",
        headers={
            "Content-Disposition": 'attachment; filename="template_produk_accurate.csv"'
        }
    )

@router.get("/template/contacts", summary="Download contoh template CSV kontak Salesforce")
async def download_contacts_template():
    """Mengunduh file template CSV siap pakai untuk import kontak."""
    csv_content = (
        "Nama,No HP,Email,Perusahaan,Tipe\n"
        "Budi Santoso,081234567890,budi.santoso@gmail.com,PT Maju Bersama,person\n"
        "Siti Rahma,085678901234,siti.rahma@techindo.co.id,TechIndo Solusi,person\n"
        "CV Surya Sparepart,082198765432,sales@suryasparepart.com,CV Surya Sparepart,org\n"
        "dr. Andi Wijaya,081122334455,andi.wijaya@medikaclinic.id,Medika Clinic,person\n"
    )
    return Response(
        content=csv_content,
        media_type="text/csv",
        headers={
            "Content-Disposition": 'attachment; filename="template_kontak_salesforce.csv"'
        }
    )
